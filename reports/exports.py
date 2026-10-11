from django.http import HttpResponse
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_FILL = PatternFill("solid", fgColor="7A1F2B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
# Cream -> light saffron -> saffron. It stops at saffron so black numbers stay readable.
SHADE_COLORS = ("FFF8EE", "F6C79C", "E8751A")


def _shade(sheet, first_row, last_row, columns):
    """One Excel colour scale over a group of columns (1-based), e.g. every Hearing column of a daily sheet.

    0 is the palest and the 90th percentile and above the darkest, so one very
    long day does not wash out the rest. Empty cells are left unshaded by Excel.
    """
    if last_row < first_row or not columns:
        return
    ranges = " ".join(
        f"{get_column_letter(column)}{first_row}:{get_column_letter(column)}{last_row}" for column in columns
    )
    low, middle, high = SHADE_COLORS
    sheet.conditional_formatting.add(
        ranges,
        ColorScaleRule(
            start_type="num", start_value=0, start_color=low,
            mid_type="percentile", mid_value=50, mid_color=middle,
            end_type="percentile", end_value=90, end_color=high,
        ),
    )


def _fill_sheet(sheet, sheet_title, intro, headers, rows, shade=(), unshaded_last_rows=0):
    sheet.title = sheet_title[:31]
    for line in intro:
        sheet.append([str(line)])
    if intro:
        sheet.append([])
    sheet.append([str(header) for header in headers])
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    for columns in shade:
        _shade(sheet, header_row + 1, header_row + len(rows) - unshaded_last_rows, columns)
    for index, header in enumerate(headers, 1):
        longest = max([len(str(header))] + [len(str(row[index - 1] or "")) for row in rows])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 50)


def workbook_response(filename, sheet_title, intro, headers, rows, extra_sheets=(), shade=(), unshaded_last_rows=0):
    """An .xlsx download: a few intro lines, then a styled header row and the data.

    `extra_sheets` adds more sheets, each as (sheet_title, intro, headers, rows).
    `shade` lists groups of 1-based column numbers; each group shares one colour
    scale. `unshaded_last_rows` keeps footer rows (a group total) out of it.
    """
    workbook = Workbook()
    _fill_sheet(workbook.active, sheet_title, intro, headers, rows, shade, unshaded_last_rows)
    for extra in extra_sheets:
        _fill_sheet(workbook.create_sheet(), *extra)

    response = HttpResponse(content_type=XLSX)
    response["Content-Disposition"] = f'attachment; filename="{slugify(filename)}.xlsx"'
    workbook.save(response)
    return response

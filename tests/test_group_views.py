from datetime import timedelta
from io import BytesIO

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from sadhana.models import EntryValue, SadhanaEntry
from sadhana.services import analytics, periods


@pytest.fixture
def today(db):
    return timezone.localdate()


@pytest.fixture
def group(make_user, field, today):
    """A coordinator with two members, one independent devotee and one inactive devotee."""
    coordinator = make_user(email="coord@example.com", name="Coord", spiritual_name="Madhava Das", is_coordinator=True)
    people = {
        "madhava": coordinator,
        "arjuna": make_user(email="arjuna@example.com", name="Arjuna", coordinator=coordinator),
        "bhima": make_user(email="bhima@example.com", name="Bhima", coordinator=coordinator),
        "solo": make_user(email="solo@example.com", name="Solo"),
        "inactive": make_user(email="gone@example.com", name="Gone", is_participating=False),
    }
    User.objects.update(date_joined=timezone.now() - timedelta(days=365))

    def log(who, days_ago, hearing=None, rounds=None):
        entry = SadhanaEntry.objects.create(user=people[who], date=today - timedelta(days=days_ago))
        if hearing is not None:
            EntryValue.objects.create(entry=entry, field=field("hearing"), value_int=hearing)
        if rounds is not None:
            EntryValue.objects.create(entry=entry, field=field("rounds"), value_int=rounds)

    log("madhava", 0, hearing=60, rounds=16)
    log("madhava", 1, hearing=30, rounds=16)
    log("arjuna", 0, hearing=90, rounds=16)
    log("bhima", 1, hearing=30, rounds=8)
    log("solo", 0, hearing=90, rounds=4)
    log("inactive", 0, hearing=600, rounds=64)
    return people


def sheet_rows(response):
    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    assert response["Content-Disposition"].endswith('.xlsx"')
    return [list(row) for row in load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True)]


# --- Analysis -------------------------------------------------------------


def test_group_pages_require_login(client, db):
    for name in ["analysis", "ranks", "analysis_export", "ranks_export"]:
        assert client.get(reverse(name)).status_code == 302


def test_totals_and_per_person_average_exclude_inactive_users(client, group):
    client.force_login(group["solo"])
    response = client.get(reverse("analysis"))
    summary = response.context["summary"]
    # 4 participating devotees; the inactive devotee's 600 minutes are not counted.
    assert (summary.people, summary.total) == (4, 300)
    assert summary.average_per_person == 75.0
    assert summary.average_per_person_per_day == round(75 / 7, 1)
    assert "Gone" not in [row.user.display_name for row in response.context["rows"]]
    html = response.content.decode()
    assert "Average per person" in html and "Per person per day" in html


def test_group_summary_with_nobody_in_scope(today, db):
    summary = analytics.group_summary([], periods.resolve("last_7", today), today)
    assert (summary.people, summary.total, summary.average_per_person, summary.average_per_person_per_day) == (0, 0, 0, 0.0)


def test_coordinator_scope_is_the_coordinator_plus_team(client, group):
    client.force_login(group["solo"])
    response = client.get(reverse("analysis"), {"coordinator": group["madhava"].pk})
    assert {row.user.display_name for row in response.context["rows"]} == {"Madhava Das", "Arjuna", "Bhima"}
    summary = response.context["summary"]
    assert (summary.people, summary.total, summary.average_per_person) == (3, 210, 70.0)


def test_per_user_table(client, group):
    client.force_login(group["solo"])
    rows = client.get(reverse("analysis")).context["rows"]
    table = {row.user.display_name: row for row in rows}
    assert [row.stats.total for row in rows] == sorted((row.stats.total for row in rows), reverse=True)
    madhava = table["Madhava Das"]
    assert (madhava.stats.total, madhava.stats.days_filled, madhava.stats.daily_average) == (90, 2, round(90 / 7, 1))
    assert (madhava.streaks.current, madhava.streaks.on_time_current) == (2, 2)
    assert table["Bhima"].streaks.current == 1  # yesterday only: the grace rule keeps it alive


def test_field_and_period_filters(client, group):
    client.force_login(group["solo"])
    response = client.get(reverse("analysis"), {"field": "rounds", "period": "last_30"})
    assert response.context["field"].key == "rounds"
    assert response.context["summary"].total == 60
    assert response.context["rows"][0].streaks is None
    assert "Streak" not in response.content.decode().split("By devotee")[1].split("</thead>")[0]
    assert [field.key for field in response.context["fields"]] == ["hearing", "reading", "seva", "rounds", "mangala_arati"]


def test_community_trend_is_two_single_series_charts(client, group):
    client.force_login(group["solo"])
    charts = client.get(reverse("analysis")).context["charts"]
    assert len(charts) == 2 and all(len(chart["datasets"]) == 1 for chart in charts)
    assert charts[0]["datasets"][0]["data"][-2:] == [60, 240]
    assert charts[1]["datasets"][0]["data"][-2:] == [15.0, 60.0]


def test_not_filled_today(client, group):
    client.force_login(group["solo"])
    response = client.get(reverse("analysis"))
    assert [person.display_name for person in response.context["not_filled"]] == ["Bhima"]
    scoped = client.get(reverse("analysis"), {"coordinator": group["madhava"].pk})
    assert [person.display_name for person in scoped.context["not_filled"]] == ["Bhima"]


def test_analysis_export(client, group):
    client.force_login(group["solo"])
    rows = sheet_rows(client.get(reverse("analysis_export"), {"coordinator": group["madhava"].pk}))
    header = rows.index(["Name", "Coordinator", "Total", "Daily average", "Days filled", "Current streak", "On-time streak"])
    assert rows[0][0] == "Hearing — Madhava Das"
    # Sorted by total; Arjuna and Madhava Das tie on 90 and are then ordered by name.
    assert rows[header + 1:] == [
        ["Arjuna", "Madhava Das", 90, round(90 / 7, 1), 1, 1, 1],
        ["Madhava Das", None, 90, round(90 / 7, 1), 2, 2, 2],
        ["Bhima", "Madhava Das", 30, round(30 / 7, 1), 1, 1, 1],
    ]


# --- Ranks ----------------------------------------------------------------


def ranked(response):
    return [(item["rank"], item["row"].user.display_name, item["value"]) for item in response.context["ranked"]]


def test_rank_by_total_with_ties(client, group):
    client.force_login(group["solo"])
    response = client.get(reverse("ranks"))
    assert response.context["option"].key == "total"
    assert ranked(response) == [(1, "Arjuna", 90), (1, "Madhava Das", 90), (1, "Solo", 90), (4, "Bhima", 30)]
    html = response.content.decode()
    assert "🥇" in html and "rank-me" in html
    assert "Gone" not in html  # only participating users


def test_rank_by_streak(client, group):
    client.force_login(group["solo"])
    response = client.get(reverse("ranks"), {"metric": "hearing", "by": "current"})
    assert ranked(response) == [(1, "Madhava Das", 2), (2, "Arjuna", 1), (2, "Bhima", 1), (2, "Solo", 1)]
    assert "no period applies" in response.content.decode()
    assert 'name="period"' not in response.content.decode().split("no period applies")[0].split("</form>")[-1]


def test_streak_options_hidden_for_rounds(client, group):
    client.force_login(group["solo"])
    response = client.get(reverse("ranks"), {"metric": "rounds", "by": "current"})
    assert [option.key for option in response.context["options"]] == ["total", "daily_average", "days_filled"]
    assert response.context["option"].key == "total"
    assert "Current streak" not in response.content.decode()
    assert ranked(response) == [(1, "Madhava Das", 32), (2, "Arjuna", 16), (3, "Bhima", 8), (4, "Solo", 4)]


def test_rank_days_filled_and_average(client, group):
    client.force_login(group["solo"])
    filled = ranked(client.get(reverse("ranks"), {"by": "days_filled"}))
    assert filled[0] == (1, "Madhava Das", 2) and [rank for rank, _, _ in filled] == [1, 2, 2, 2]
    average = ranked(client.get(reverse("ranks"), {"by": "daily_average", "period": "last_10"}))
    assert average[0][2] == 9.0 and average[-1] == (4, "Bhima", 3.0)


def test_rank_coordinator_filter_and_period(client, group):
    client.force_login(group["solo"])
    response = client.get(reverse("ranks"), {"coordinator": group["madhava"].pk})
    assert [name for _, name, _ in ranked(response)] == ["Arjuna", "Madhava Das", "Bhima"]
    last_month = client.get(reverse("ranks"), {"period": "last_month"})
    assert last_month.context["period"].key == "last_month"


def test_ranks_export(client, group):
    client.force_login(group["solo"])
    rows = sheet_rows(client.get(reverse("ranks_export"), {"metric": "rounds"}))
    header = rows.index(["Rank", "Name", "Coordinator", "Total"])
    assert rows[0][0] == "Rounds — Total"
    assert rows[header + 1:] == [
        [1, "Madhava Das", None, 32], [2, "Arjuna", "Madhava Das", 16], [3, "Bhima", "Madhava Das", 8], [4, "Solo", None, 4],
    ]


def test_dashboard_export(client, group, today):
    client.force_login(group["solo"])
    rows = sheet_rows(client.get(reverse("dashboard_export", args=[group["madhava"].pk])))
    assert rows[0][0] == "Madhava Das"
    header = rows.index(["Date", "Hearing", "Book Reading", "Seva", "Rounds", "Notes / Realizations", "On time"])
    first = rows[header + 1]
    assert first[0].date() == today and first[1:] == [60, None, None, 16, None, "Yes"]
    assert len(rows[header + 1:]) == 2
    # The page links to the export with the current filters.
    page = client.get(reverse("dashboard", args=[group["madhava"].pk]), {"period": "last_30"}).content.decode()
    assert f"{reverse('dashboard_export', args=[group['madhava'].pk])}?period=last_30" in page

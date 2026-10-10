from datetime import date, timedelta
from io import BytesIO

import pytest
from django.db import IntegrityError, transaction
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from audit.models import AuditLog
from sadhana.models import EntryValue, Festival, SadhanaEntry
from sadhana.services import analytics, festivals, periods


@pytest.fixture
def today(db):
    return timezone.localdate()


def festival(name, start, end=None):
    return Festival(name=name, start_date=start, end_date=end or start)


# --- Which dates are festival days ----------------------------------------


def test_single_day_range_and_overlap():
    found = festivals.festival_dates(
        date(2026, 8, 1),
        date(2026, 8, 31),
        [
            festival("Janmashtami", date(2026, 8, 14)),
            festival("Jhulan Yatra", date(2026, 8, 10), date(2026, 8, 14)),
            festival("Next month", date(2026, 9, 5)),
        ],
    )
    assert sorted(found) == [date(2026, 8, day) for day in range(10, 15)]
    assert found[date(2026, 8, 12)] == ["Jhulan Yatra"]
    # Overlap: the date belongs to both.
    assert found[date(2026, 8, 14)] == ["Jhulan Yatra", "Janmashtami"]


def test_ranges_are_clipped_to_the_window():
    kartik = festival("Kartik", date(2026, 10, 25), date(2026, 11, 24))
    found = festivals.festival_dates(date(2026, 11, 1), date(2026, 11, 3), [kartik])
    assert sorted(found) == [date(2026, 11, 1), date(2026, 11, 2), date(2026, 11, 3)]
    assert festivals.festival_dates(date(2026, 12, 1), date(2026, 12, 31), [kartik]) == {}


def test_festival_dates_from_the_database(db):
    Festival.objects.create(name="Gaura Purnima", start_date=date(2026, 3, 3), end_date=date(2026, 3, 3))
    assert festivals.festival_dates(date(2026, 3, 1), date(2026, 3, 31)) == {date(2026, 3, 3): ["Gaura Purnima"]}


def test_end_date_cannot_be_before_start_date(db):
    with pytest.raises(IntegrityError), transaction.atomic():
        Festival.objects.create(name="Backwards", start_date=date(2026, 3, 3), end_date=date(2026, 3, 1))


def test_day_filter():
    days = {date(2026, 8, 14): ["Janmashtami"]}
    assert festivals.day_filter("all", days) is None
    assert festivals.day_filter("regular", days)(date(2026, 8, 14)) is False
    assert festivals.day_filter("regular", days)(date(2026, 8, 15)) is True
    assert festivals.day_filter("festival", days)(date(2026, 8, 14)) is True
    assert festivals.choose("nonsense") == "all" and festivals.choose("regular") == "regular"


# --- Effect on the numbers ------------------------------------------------


@pytest.fixture
def festival_week(make_user, field, today):
    """Two devotees with ten steady days of 60 minutes seva; four days ago was a festival (300 and 240)."""
    people = [make_user(email="a@example.com", name="Arjuna"), make_user(email="b@example.com", name="Bhima")]
    User.objects.update(date_joined=timezone.now() - timedelta(days=365))
    for person, festival_minutes in zip(people, (300, 240)):
        for days_ago in range(10):
            entry = SadhanaEntry.objects.create(user=person, date=today - timedelta(days=days_ago))
            minutes = festival_minutes if days_ago == 4 else 60
            EntryValue.objects.create(entry=entry, field=field("seva"), value_int=minutes)
            EntryValue.objects.create(entry=entry, field=field("hearing"), value_int=30)
    Festival.objects.create(name="Janmashtami", start_date=today - timedelta(days=4), end_date=today - timedelta(days=4))
    return people


def test_regular_days_only_changes_totals_and_averages_but_not_streaks(festival_week, field, today):
    arjuna = festival_week[0]
    days = analytics.load_days(arjuna)
    period = periods.resolve("last_10", today)
    festival_days = festivals.festival_dates(period.start, period.end)
    seva = field("seva")

    everything = analytics.field_stats(days, seva, period, arjuna.date_joined, today)
    assert (everything.total, everything.daily_average, everything.days_filled) == (840, 84.0, 10)
    # The typical day (median) is not moved by the one long day; the average is.
    assert everything.typical_day == 60

    regular = analytics.field_stats(days, seva, period, arjuna.date_joined, today, festivals.day_filter("regular", festival_days))
    assert (regular.total, regular.daily_average, regular.days_filled, regular.typical_day) == (540, 60.0, 9, 60)

    only = analytics.field_stats(days, seva, period, arjuna.date_joined, today, festivals.day_filter("festival", festival_days))
    assert (only.total, only.daily_average, only.days_filled, only.typical_day) == (300, 300.0, 1, 300)

    rows = analytics.user_rows([arjuna], {arjuna.pk: days}, seva, period, today, festivals.day_filter("regular", festival_days))
    assert rows[0].streaks.current == 10 and rows[0].streaks.longest == 10


def test_analysis_day_filter_and_festival_summary(client, festival_week):
    client.force_login(festival_week[0])
    url = reverse("analysis")
    base = {"field": "seva", "period": "last_10"}

    everything = client.get(url, base)
    assert everything.context["summary"].total == 1620
    assert everything.context["days_key"] == "all"
    summary = everything.context["festival_summaries"]
    assert [(item["festival"].name, item["total"], item["devotees"]) for item in summary] == [("Janmashtami", 540, 2)]
    html = everything.content.decode()
    assert "Janmashtami" in html and "across 2 devotees" in html and "Regular days only" in html

    regular = client.get(url, {**base, "days": "regular"})
    assert regular.context["summary"].total == 1080
    assert regular.context["summary"].average_per_person == 540.0
    assert regular.context["summary"].average_per_person_per_day == 60.0
    assert [row.streaks.current for row in regular.context["rows"]] == [10, 10]
    # The filter survives a change of period.
    assert 'name="days" value="regular"' in regular.content.decode()

    only = client.get(url, {**base, "days": "festival"})
    assert only.context["summary"].total == 540
    # No festival in the period: no summary.
    last_year = {"field": "seva", "period": "custom", "start": "2025-01-01", "end": "2025-01-31"}
    assert client.get(url, last_year).context["festival_summaries"] == []


def test_rank_day_filter_applies_to_totals_but_not_streaks(client, festival_week):
    client.force_login(festival_week[0])
    url = reverse("ranks")

    def values(**params):
        response = client.get(url, {"metric": "seva", "period": "last_10", **params})
        return {item["row"].user.display_name: item["value"] for item in response.context["ranked"]}

    assert values() == {"Arjuna": 840, "Bhima": 780}
    assert values(days="regular") == {"Arjuna": 540, "Bhima": 540}
    assert values(days="festival") == {"Arjuna": 300, "Bhima": 240}
    assert values(days="regular", by="daily_average") == {"Arjuna": 60.0, "Bhima": 60.0}
    assert values(days="festival", by="current") == {"Arjuna": 10, "Bhima": 10}


# --- Charts and heatmap ---------------------------------------------------


def test_chart_is_capped_but_keeps_real_values(client, festival_week, today):
    arjuna = festival_week[0]
    client.force_login(arjuna)
    response = client.get(reverse("dashboard", args=[arjuna.pk]), {"period": "last_10"})
    seva = next(chart for chart in response.context["charts"] if chart["title"] == "Seva")
    # Usual range: 90th percentile (84) × 1.2 = 100.8 → 120. The festival day is clipped, not changed.
    assert seva["cap"] == 120
    assert seva["values"][5] == 300
    assert seva["clipped"] == [{"index": 5, "label": (today - timedelta(days=4)).strftime("%-d %b"), "value": 300}]
    assert seva["festivals"] == {5: "Janmashtami"}
    hearing = next(chart for chart in response.context["charts"] if chart["title"] == "Hearing")
    assert hearing["cap"] is None and hearing["festivals"] == {5: "Janmashtami"}
    html = response.content.decode()
    assert "Usual range" in html and "Full range" in html
    assert "Above the usual range:" in html and "<strong>300</strong>" in html


def test_heatmap_is_coloured_by_the_usual_range(client, festival_week, today):
    arjuna = festival_week[0]
    client.force_login(arjuna)
    response = client.get(reverse("dashboard", args=[arjuna.pk]), {"heat": "seva"})
    assert response.context["heat_top"] == 120
    cells = {cell["date"]: cell for column in response.context["heatmap"] for cell in column["cells"] if cell}
    # A normal 60-minute day is mid-scale instead of the palest colour; the festival day is the darkest.
    assert cells[today]["level"] == 2
    assert cells[today - timedelta(days=4)]["level"] == 4
    assert cells[today - timedelta(days=4)]["value"] == 300
    assert "300 min" in response.content.decode()


def test_heatmap_without_a_cap_scales_to_the_largest_day(festival_week, field, today):
    days = analytics.load_days(festival_week[0])
    cells = {cell["date"]: cell for column in analytics.heatmap(days, field("seva"), today) for cell in column["cells"] if cell}
    assert cells[today]["level"] == 1 and cells[today - timedelta(days=4)]["level"] == 4


def test_exports_contain_real_unclipped_values(client, festival_week, today):
    arjuna = festival_week[0]
    client.force_login(arjuna)

    def rows(name, *args, **params):
        response = client.get(reverse(name, args=args), params)
        return [list(row) for row in load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True)]

    entries = rows("dashboard_export", arjuna.pk, period="last_10")
    header = next(index for index, row in enumerate(entries) if row[0] == "Date")
    seva_column = entries[header].index("Seva")
    assert max(row[seva_column] for row in entries[header + 1:]) == 300

    analysis = rows("analysis_export", field="seva", period="last_10")
    # Columns: name, coordinator, total, daily average, days filled, current streak, on-time streak.
    # (The fixture's older entries were all created today, so only two days count as on time.)
    assert ["Arjuna", None, 840, 84.0, 10, 10, 2] in analysis
    regular = rows("analysis_export", field="seva", period="last_10", days="regular")
    assert ["Arjuna", None, 540, 60.0, 9, 10, 2] in regular
    assert any(row[0] == "Regular days only" for row in regular)

    ranks = rows("ranks_export", metric="seva", period="last_10")
    assert [1, "Arjuna", None, 840] in ranks


# --- Manage page ----------------------------------------------------------


def test_manage_festivals(client, superuser, user):
    client.force_login(user)
    assert client.get(reverse("manage_festivals")).status_code == 403
    assert client.post(reverse("manage_festival_add"), {}).status_code == 403

    client.force_login(superuser)
    assert client.get(reverse("manage_festivals")).status_code == 200
    data = {"name": "Janmashtami", "start_date": "2026-08-14", "end_date": "2026-08-14", "notes": ""}
    assert client.post(reverse("manage_festival_add"), data).status_code == 302
    created = Festival.objects.get()
    assert AuditLog.objects.get(action="festival_create").new_value == "Janmashtami (2026-08-14 to 2026-08-14)"

    bad = client.post(reverse("manage_festival_edit", args=[created.pk]), {**data, "end_date": "2026-08-10"})
    assert bad.status_code == 200 and "end_date" in bad.context["form"].errors

    client.post(reverse("manage_festival_edit", args=[created.pk]), {**data, "end_date": "2026-08-15"})
    created.refresh_from_db()
    assert created.end_date == date(2026, 8, 15)
    update = AuditLog.objects.get(action="festival_update")
    assert (update.field_key, update.old_value, update.new_value) == (
        f"festival:{created.pk}", "end_date=2026-08-14", "end_date=2026-08-15",
    )
    assert "Janmashtami" in client.get(reverse("manage_festivals")).content.decode()

    assert client.get(reverse("manage_festival_delete", args=[created.pk])).status_code == 405
    client.post(reverse("manage_festival_delete", args=[created.pk]))
    assert not Festival.objects.exists()
    assert AuditLog.objects.filter(action="festival_delete").count() == 1


def test_field_target_is_seeded_and_editable(client, superuser, field):
    assert field("rounds").target == 16 and field("hearing").target is None
    client.force_login(superuser)
    assert 'name="target"' in client.get(reverse("manage_field_edit", args=[field("rounds").pk])).content.decode()

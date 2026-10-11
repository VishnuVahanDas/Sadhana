"""Simple / Advanced display mode, and Analysis for superusers only (docs/UPDATE-2026-10-simple-advanced-mode.md)."""

from datetime import timedelta
from io import BytesIO

import pytest
from django.contrib.messages import get_messages
from django.core.management import call_command
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from audit.models import AuditLog
from sadhana.models import AppSettings, EntryValue, SadhanaEntry, SadhanaField
from sadhana.services import streaks


@pytest.fixture
def today(db):
    return timezone.localdate()


@pytest.fixture
def people(make_user, today):
    """A Simple and an Advanced devotee, each with three days of hearing, reading, seva and rounds."""
    simple = make_user(email="simple@example.com", name="Sarala")
    advanced = make_user(email="advanced@example.com", name="Vistara", ui_mode="advanced")
    fields = {field.key: field for field in SadhanaField.objects.all()}
    for person in (simple, advanced):
        for days_ago in range(3):
            entry = SadhanaEntry.objects.create(user=person, date=today - timedelta(days=days_ago))
            for key, value in (("hearing", 30), ("reading", 20), ("seva", 40), ("rounds", 16)):
                EntryValue.objects.create(entry=entry, field=fields[key], value_int=value)
    return simple, advanced


@pytest.fixture
def no_streaks(monkeypatch):
    """Fail the test if anything works out a streak."""

    def forbidden(*args, **kwargs):
        raise AssertionError("streaks were computed for a Simple-mode reader")

    monkeypatch.setattr(streaks, "compute_streaks", forbidden)


def streak_pages(reader, other, today):
    return [
        reverse("home"),
        reverse("entry", args=[today]),
        reverse("dashboard", args=[reader.pk]),
        reverse("dashboard", args=[other.pk]),
        reverse("dashboard", args=[other.pk]) + "?compare=1",
        reverse("yatra"),
    ]


# --- The field ------------------------------------------------------------


def test_new_users_default_to_simple(make_user, superuser):
    assert make_user().ui_mode == "simple"
    assert superuser.ui_mode == "simple" and not superuser.is_advanced
    assert User.objects.create_superuser(email="root@example.com", password="x", name="Root").ui_mode == "simple"
    assert make_user(email="adv@example.com", ui_mode="advanced").is_advanced


@pytest.mark.django_db(transaction=True)
def test_migration_sets_existing_users_to_simple_and_keeps_a_database_default():
    executor = MigrationExecutor(connection)
    try:
        executor.migrate([("accounts", "0002_user_theme")])
        Old = executor.loader.project_state([("accounts", "0002_user_theme")]).apps.get_model("accounts", "User")
        Old.objects.create(email="old@example.com", name="Old Devotee", password="!")
        assert not hasattr(Old, "ui_mode")

        executor = MigrationExecutor(connection)
        executor.migrate([("accounts", "0003_user_ui_mode")])
        assert User.objects.get(email="old@example.com").ui_mode == "simple"
        # Rollback safety: the previous release does not know the column, and its inserts must still work.
        Old.objects.create(email="older-code@example.com", name="Added By Old Code", password="!")
        assert User.objects.get(email="older-code@example.com").ui_mode == "simple"
    finally:
        executor = MigrationExecutor(connection)
        executor.migrate(executor.loader.graph.leaf_nodes())
        call_command("flush", verbosity=0, interactive=False)


# --- Analysis: superusers only ----------------------------------------------


@pytest.mark.parametrize("name", ["analysis", "analysis_export"])
def test_analysis_redirects_a_normal_user_to_today(client, make_user, name):
    for mode in ("simple", "advanced"):
        client.force_login(make_user(email=f"{mode}@example.com", ui_mode=mode))
        response = client.get(reverse(name), {"field": "hearing", "compare": "1"})
        assert response.status_code == 302 and response.url == reverse("home")
        assert "Analysis is for the temple coordinators" in str(list(get_messages(response.wsgi_request))[-1])
    assert client.get(reverse("home")).status_code == 200  # a friendly note, not an error page


def test_analysis_is_open_to_a_superuser_in_either_mode(client, superuser, people):
    for mode in ("simple", "advanced"):
        superuser.ui_mode = mode
        superuser.save()
        client.force_login(superuser)
        response = client.get(reverse("analysis"), {"field": "hearing"})
        assert response.status_code == 200
        # Analysis keeps its streak columns whatever the superuser's own mode is.
        assert "Dhairya" in response.content.decode()
        assert all(row.streaks.current == 3 for row in response.context["rows"] if row.user != superuser)
        export = client.get(reverse("analysis_export"), {"field": "hearing"})
        assert export.status_code == 200
        header = next(row for row in load_workbook(BytesIO(export.content)).active.values if row[0] == "Name")
        assert "Dhairya (current streak)" in header and "Dhairya (on-time streak)" in header


def test_analysis_link_is_in_the_navigation_for_superusers_only(client, make_user, superuser):
    for mode in ("simple", "advanced"):
        client.force_login(make_user(email=f"{mode}@example.com", ui_mode=mode))
        for page in (reverse("home"), reverse("sheet"), reverse("yatra")):
            html = client.get(page).content.decode()
            assert reverse("analysis") not in html and "Back to Analysis" not in html
    client.force_login(superuser)
    assert reverse("analysis") in client.get(reverse("home")).content.decode()
    assert "Back to Analysis" in client.get(reverse("sheet")).content.decode()


def test_sheet_and_devotees_stay_open_to_everyone(client, people):
    simple, _advanced = people
    client.force_login(simple)
    for name in ("sheet", "sheet_export", "devotees", "yatra"):
        assert client.get(reverse(name)).status_code == 200


# --- Streaks follow the reader's mode ----------------------------------------


def test_simple_reader_sees_no_streaks_and_none_are_computed(client, people, today, no_streaks):
    simple, advanced = people
    client.force_login(simple)
    for url in streak_pages(simple, advanced, today):
        response = client.get(url)
        assert response.status_code == 200, url
        html = response.content.decode()
        # (The Display mode switch itself says "streaks", and the verse on Sanga Yatra names dhairya.)
        for text in ("streak-tile", "streak-number", "🔥", "Longest Dhairya", "On-time Dhairya", "Dhairya —"):
            assert text not in html, (url, text)
        if url != reverse("yatra"):
            assert "Dhairya" not in html, url
        assert response.context["show_streaks"] is False

    assert client.get(reverse("home")).context["streak_badges"] == []
    report = client.get(reverse("dashboard", args=[advanced.pk]))
    assert all(card["streaks"] is None for card in report.context["cards"])
    assert report.context["journey"].dhairya == [] and report.context["journey"].together is None
    road = client.get(reverse("yatra")).context
    assert road["longest_dhairya"] is None
    assert all(item["flame"] is None for group in road["groups"] for item in group["people"])


def test_simple_reader_keeps_everything_else(client, people, today):
    simple, _advanced = people
    client.force_login(simple)
    today_page = client.get(reverse("entry", args=[today])).content.decode()
    assert 'aria-label="Bhakti Yatra"' in today_page and "min offered" in today_page  # vehicle and Utsāha stay
    report = client.get(reverse("dashboard", args=[simple.pk]))
    html = report.content.decode()
    for text in ("Daily average", "Typical day", "Days filled", 'class="heatmap"'):
        assert text in html, text
    assert report.context["charts"] and report.context["heatmap"]
    assert client.get(reverse("dashboard_export", args=[simple.pk])).status_code == 200


def test_advanced_reader_sees_streaks_on_own_and_other_pages(client, people, today):
    simple, advanced = people
    client.force_login(advanced)
    home = client.get(reverse("entry", args=[today]))
    assert {badge["field"].key: badge["streaks"].current for badge in home.context["streak_badges"]} == {
        "hearing": 3, "reading": 3, "seva": 3, "rounds": 3,
    }
    assert "Dhairya" in home.content.decode()
    # The mode belongs to the reader: an Advanced reader sees a Simple devotee's streaks too.
    report = client.get(reverse("dashboard", args=[simple.pk]))
    streak_cards = {card["field"].key: card["streaks"].current for card in report.context["cards"] if card["streaks"]}
    assert streak_cards == {"hearing": 3, "reading": 3, "seva": 3, "rounds": 3}
    assert "On-time Dhairya" in report.content.decode()
    assert [item["field"].key for item in report.context["journey"].dhairya] == ["hearing", "reading", "seva"]
    road = client.get(reverse("yatra"))
    assert road.context["longest_dhairya"]["streaks"].longest == 3
    assert all(item["flame"] == 3 for group in road.context["groups"] for item in group["people"])
    assert "🔥 3" in road.content.decode()


def test_superuser_follows_their_own_mode(client, superuser, people, today, no_streaks):
    """A superuser in Simple mode has no streaks on Today or My Report, but still has Analysis in the bar."""
    simple, _advanced = people
    client.force_login(superuser)
    # A superuser editing someone else's day sees it in their own (Simple) mode as well.
    for url in (reverse("home"), reverse("dashboard", args=[superuser.pk]), reverse("entry", args=[today]) + f"?user={simple.pk}"):
        html = client.get(url).content.decode()
        assert "Dhairya" not in html and "🔥" not in html
        assert reverse("analysis") in html


def test_rank_table_has_no_streak_options_for_a_simple_reader(client, people, no_streaks):
    settings = AppSettings.load()
    settings.rank_table_enabled = True
    settings.save()
    simple, _advanced = people
    client.force_login(simple)
    response = client.get(reverse("ranks"), {"metric": "hearing", "by": "current"})  # typed in by hand
    assert response.context["option"].key == "total"
    assert [option.key for option in response.context["options"]] == ["total", "daily_average", "days_filled"]
    assert "Dhairya" not in response.content.decode()
    export = client.get(reverse("ranks_export"), {"metric": "hearing", "by": "longest"})
    header = next(row for row in load_workbook(BytesIO(export.content)).active.values if row[0] == "Rank")
    assert header[3] == "Total"


def test_rank_table_keeps_streak_options_for_an_advanced_reader(client, people):
    settings = AppSettings.load()
    settings.rank_table_enabled = True
    settings.save()
    _simple, advanced = people
    client.force_login(advanced)
    response = client.get(reverse("ranks"), {"metric": "hearing", "by": "current"})
    assert response.context["option"].key == "current" and len(response.context["options"]) == 7
    assert [item["value"] for item in response.context["ranked"]] == [3, 3]


def test_individual_excel_has_no_streak_columns_in_either_mode(client, people):
    simple, advanced = people
    for reader in (simple, advanced):
        client.force_login(reader)
        sheet = load_workbook(BytesIO(client.get(reverse("dashboard_export", args=[simple.pk])).content)).active
        header = next(row for row in sheet.values if row[0] == "Date")
        assert header[-1] == "On time" and not any("streak" in str(cell).lower() or "Dhairya" in str(cell) for cell in header)


# --- Changing the mode --------------------------------------------------------


def test_devotee_switches_their_own_mode(client, user):
    client.force_login(user)
    assert 'action="%s"' % reverse("my_mode") in client.get(reverse("home")).content.decode()
    response = client.post(reverse("my_mode"), {"ui_mode": "advanced", "next": reverse("me")})
    assert response.status_code == 302 and response.url == reverse("me")
    user.refresh_from_db()
    assert user.is_advanced
    log = AuditLog.objects.get(action="ui_mode_change")
    assert (log.actor, log.target_user, log.field_key, log.old_value, log.new_value) == (
        user, user, "ui_mode", "simple", "advanced",
    )
    assert "Dhairya" in client.get(reverse("home")).content.decode()

    # Same mode again, or nonsense: nothing changes and nothing is logged.
    client.post(reverse("my_mode"), {"ui_mode": "advanced"})
    client.post(reverse("my_mode"), {"ui_mode": "expert"})
    user.refresh_from_db()
    assert user.is_advanced and AuditLog.objects.filter(action="ui_mode_change").count() == 1

    client.post(reverse("my_mode"), {"ui_mode": "simple", "next": "https://evil.example/"})
    user.refresh_from_db()
    assert user.ui_mode == "simple" and AuditLog.objects.filter(action="ui_mode_change").count() == 2
    assert client.get(reverse("my_mode")).status_code == 405


def test_superuser_sets_one_users_mode_on_the_edit_form(client, superuser, user):
    client.force_login(superuser)
    data = {
        "email": user.email, "name": user.name, "spiritual_name": "", "phone": "", "coordinator": "",
        "is_participating": "on", "ui_mode": "advanced",
    }
    assert client.post(reverse("manage_user_edit", args=[user.pk]), data).status_code == 302
    user.refresh_from_db()
    assert user.is_advanced
    log = AuditLog.objects.get(action="ui_mode_change")
    assert (log.actor, log.target_user, log.old_value, log.new_value) == (superuser, user, "simple", "advanced")
    assert not AuditLog.objects.filter(action="user_update").exists()


def test_superuser_bulk_sets_modes_with_one_audit_row_per_changed_user(client, superuser, make_user):
    a = make_user(email="a@example.com", name="Arjuna")
    b = make_user(email="b@example.com", name="Bhima")
    already = make_user(email="c@example.com", name="Chaitanya", ui_mode="advanced")
    untouched = make_user(email="d@example.com", name="Damodara")
    client.force_login(superuser)
    response = client.post(
        reverse("manage_users_mode"), {"ui_mode": "advanced", "users": [a.pk, b.pk, already.pk], "filter": "simple"}
    )
    assert response.status_code == 302 and response.url == reverse("manage_users") + "?mode=simple"
    modes = dict(User.objects.values_list("email", "ui_mode"))
    assert modes["a@example.com"] == modes["b@example.com"] == modes["c@example.com"] == "advanced"
    assert modes["d@example.com"] == "simple" and untouched.ui_mode == "simple"
    logs = AuditLog.objects.filter(action="ui_mode_change")
    assert {log.target_user for log in logs.all()} == {a, b}  # Chaitanya was already Advanced
    assert all((log.actor, log.old_value, log.new_value) == (superuser, "simple", "advanced") for log in logs.all())

    client.post(reverse("manage_users_mode"), {"ui_mode": "simple", "users": [a.pk]})
    assert User.objects.get(pk=a.pk).ui_mode == "simple" and logs.count() == 3
    # Nothing ticked, or no valid mode: nothing changes.
    client.post(reverse("manage_users_mode"), {"ui_mode": "simple"})
    client.post(reverse("manage_users_mode"), {"ui_mode": "expert", "users": [b.pk]})
    assert User.objects.get(pk=b.pk).ui_mode == "advanced" and logs.count() == 3


def test_bulk_mode_is_superuser_only_and_post_only(client, user, superuser):
    client.force_login(user)
    assert client.post(reverse("manage_users_mode"), {"ui_mode": "advanced", "users": [user.pk]}).status_code == 403
    user.refresh_from_db()
    assert user.ui_mode == "simple"
    client.force_login(superuser)
    assert client.get(reverse("manage_users_mode")).status_code == 405


def test_user_list_shows_and_filters_by_mode(client, superuser, make_user):
    make_user(email="a@example.com", name="Arjuna", ui_mode="advanced")
    make_user(email="b@example.com", name="Bhima")
    client.force_login(superuser)
    html = client.get(reverse("manage_users")).content.decode()
    assert "Set to Simple" in html and "Set to Advanced" in html and 'name="users"' in html
    assert [person.name for person in client.get(reverse("manage_users"), {"mode": "advanced"}).context["users"]] == ["Arjuna"]
    assert {person.name for person in client.get(reverse("manage_users"), {"mode": "simple"}).context["users"]} == {"Admin", "Bhima"}
    assert client.get(reverse("manage_users"), {"mode": "nonsense"}).context["users"].count() == 3

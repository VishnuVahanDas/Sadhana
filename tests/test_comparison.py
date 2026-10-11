from datetime import date, timedelta
from io import BytesIO

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from sadhana.models import EntryValue, Festival, SadhanaEntry
from sadhana.services import comparison, periods
from sadhana.services.analytics import FieldStats
from sadhana.services.comparison import Change

SATURDAY = date(2026, 10, 10)


def previous(key, today=SATURDAY, start=None, end=None):
    earlier = periods.previous(periods.resolve(key, today, start, end))
    return earlier.start, earlier.end


# --- The previous period --------------------------------------------------


@pytest.mark.parametrize(
    "key, start, end",
    [
        ("last_7", date(2026, 9, 27), date(2026, 10, 3)),
        ("last_10", date(2026, 9, 21), date(2026, 9, 30)),
        ("last_15", date(2026, 9, 11), date(2026, 9, 25)),
        ("last_30", date(2026, 8, 12), date(2026, 9, 10)),
        ("this_week", date(2026, 9, 28), date(2026, 10, 3)),  # Monday to the same weekday
        ("last_week", date(2026, 9, 21), date(2026, 9, 27)),
        ("this_month", date(2026, 9, 1), date(2026, 9, 10)),  # like with like: 1–10 Oct against 1–10 Sep
        ("last_month", date(2026, 8, 1), date(2026, 8, 31)),
    ],
)
def test_previous_period(key, start, end):
    assert previous(key) == (start, end)


def test_rolling_periods_are_back_to_back_and_equally_long():
    for key in ("last_7", "last_10", "last_15", "last_30"):
        current = periods.resolve(key, SATURDAY)
        earlier = periods.previous(current)
        assert earlier.days == current.days
        assert earlier.end + timedelta(days=1) == current.start


def test_months_of_different_lengths():
    # Whole months: October (31 days) against September (30 days).
    assert previous("last_month", date(2026, 11, 15)) == (date(2026, 9, 1), date(2026, 9, 30))
    # Month to date on the 31st: held back to the last day of the shorter month.
    assert previous("this_month", date(2026, 3, 31)) == (date(2026, 2, 1), date(2026, 2, 28))
    assert previous("this_month", date(2028, 3, 31)) == (date(2028, 2, 1), date(2028, 2, 29))
    # Across a year boundary.
    assert previous("this_month", date(2027, 1, 5)) == (date(2026, 12, 1), date(2026, 12, 5))
    assert previous("last_month", date(2027, 1, 5)) == (date(2026, 11, 1), date(2026, 11, 30))


def test_custom_range_goes_back_by_its_own_length():
    assert previous("custom", start=date(2026, 8, 11), end=date(2026, 8, 20)) == (date(2026, 8, 1), date(2026, 8, 10))


# --- Change ---------------------------------------------------------------


def test_change_up_down_same():
    up = Change(90, 60)
    assert (up.difference, up.amount, up.direction, up.percent) == (30, 30, "up", 50)
    down = Change(45.5, 60)
    assert (down.difference, down.amount, down.direction, down.percent) == (-14.5, 14.5, "down", -24)
    same = Change(60, 60)
    assert (same.difference, same.direction, same.percent) == (0, "same", 0)


def test_change_from_zero_has_no_percentage():
    change = Change(40, 0)
    assert (change.difference, change.direction, change.percent) == (40, "up", None)
    assert Change(0, 0).direction == "same"


def test_nothing_to_compare_with():
    change = Change(40, None)
    assert (change.difference, change.amount, change.direction, change.percent) == (None, None, None, None)


def test_compare_stats_leads_with_daily_average():
    metrics = comparison.compare_stats(FieldStats(700, 100.0, 7, 90), FieldStats(350, 50.0, 5, 45))
    assert [metric["key"] for metric in metrics] == ["daily_average", "typical_day", "total", "days_filled"]
    assert [metric["change"].percent for metric in metrics] == [100, 100, 100, 40]
    assert all(metric["change"].previous is None for metric in comparison.compare_stats(FieldStats(700, 100.0, 7, 90), None))


def test_align():
    assert comparison.align([1, 2, 3], 5) == [1, 2, 3, None, None]
    assert comparison.align([1, 2, 3], 2) == [1, 2]


# --- Pages ----------------------------------------------------------------


@pytest.fixture
def today(db):
    return timezone.localdate()


@pytest.fixture
def two_weeks(make_user, field, today):
    """Arjuna: 60 minutes of hearing a day this week, 30 a day the week before. Bhima joined two days ago."""
    # Analysis is for superusers.
    arjuna = make_user(email="a@example.com", name="Arjuna", is_superuser=True)
    bhima = make_user(email="b@example.com", name="Bhima")
    User.objects.filter(pk=arjuna.pk).update(date_joined=timezone.now() - timedelta(days=365))
    User.objects.filter(pk=bhima.pk).update(date_joined=timezone.now() - timedelta(days=2))
    for days_ago in range(14):
        entry = SadhanaEntry.objects.create(user=arjuna, date=today - timedelta(days=days_ago))
        EntryValue.objects.create(entry=entry, field=field("hearing"), value_int=60 if days_ago < 7 else 30)
    for days_ago in range(3):
        entry = SadhanaEntry.objects.create(user=bhima, date=today - timedelta(days=days_ago))
        EntryValue.objects.create(entry=entry, field=field("hearing"), value_int=40)
    arjuna.refresh_from_db()
    bhima.refresh_from_db()
    return arjuna, bhima


def test_comparison_is_off_by_default(client, two_weeks):
    arjuna, _ = two_weeks
    client.force_login(arjuna)
    response = client.get(reverse("dashboard", args=[arjuna.pk]))
    assert response.context["compare"] is None
    html = response.content.decode()
    assert "Compare with previous period" in html and "compare=1" in html
    assert "Compared with the previous period" not in html


def test_dashboard_comparison(client, two_weeks, today):
    arjuna, _ = two_weeks
    client.force_login(arjuna)
    response = client.get(reverse("dashboard", args=[arjuna.pk]), {"compare": "1"})
    compare = response.context["compare"]
    assert (compare["previous"].start, compare["previous"].end) == (today - timedelta(days=13), today - timedelta(days=7))
    hearing = next(item for item in compare["fields"] if item["field"].key == "hearing")
    changes = {metric["key"]: metric["change"] for metric in hearing["metrics"]}
    assert (changes["daily_average"].current, changes["daily_average"].previous, changes["daily_average"].percent) == (60.0, 30.0, 100)
    assert (changes["total"].current, changes["total"].previous) == (420, 210)
    assert changes["days_filled"].direction == "same"
    # The chart overlays the two weeks day by day.
    assert hearing["chart"]["values"] == [60] * 7 and hearing["chart"]["previous"] == [30] * 7
    assert hearing["chart"]["previous_labels"][0] == (today - timedelta(days=13)).strftime("%-d %b")
    html = response.content.decode()
    assert "▲" in html and "more" in html and "(+100%)" in html
    assert "Compared with the previous period" in html
    # The switch stays on when the period changes, and the link now turns it off.
    assert 'name="compare" value="1"' in html
    assert response.context["compare_url"] == reverse("dashboard", args=[arjuna.pk])
    # Every chart on the page has its own slot in the one JSON list.
    ids = [chart["id"] for chart in response.context["chart_data"]]
    assert ids == list(range(len(ids))) and len(ids) == html.count("data-chart=")


def test_new_devotee_has_nothing_to_compare_with(client, two_weeks):
    _, bhima = two_weeks
    client.force_login(bhima)
    response = client.get(reverse("dashboard", args=[bhima.pk]), {"compare": "1"})
    compare = response.context["compare"]
    assert compare["had_started"] is False
    hearing = next(item for item in compare["fields"] if item["field"].key == "hearing")
    assert all(metric["change"].previous is None for metric in hearing["metrics"])
    assert hearing["chart"] is None
    assert "had not started yet" in response.content.decode()


def test_analysis_comparison(client, two_weeks, today):
    arjuna, bhima = two_weeks
    client.force_login(arjuna)
    response = client.get(reverse("analysis"), {"compare": "1"})
    compare = response.context["compare"]
    summary = {str(item["label"]): item["change"] for item in compare["summary"]}
    # This week: 420 + 120 = 540; the week before: 210.
    assert (summary["Total"].current, summary["Total"].previous, summary["Total"].direction) == (540, 210, "up")
    assert summary["Average per person"].current == 270.0 and summary["Average per person"].previous == 105.0
    table = {row.user.display_name: change for row, change, *_levels in response.context["table"]}
    assert (table["Arjuna"].current, table["Arjuna"].previous, table["Arjuna"].percent) == (60.0, 30.0, 100)
    # Bhima had not joined: "—", not a drop from zero.
    assert table["Bhima"].previous is None and table["Bhima"].direction is None
    assert compare["chart"]["previous"] == [15.0] * 7
    html = response.content.decode()
    assert "Previous avg" in html and "Compared with the previous period" in html
    assert client.get(reverse("analysis")).context["compare"] is None


def test_festival_note_when_only_one_period_has_a_festival(client, two_weeks, today):
    arjuna, _ = two_weeks
    Festival.objects.create(name="Janmashtami", start_date=today - timedelta(days=2), end_date=today - timedelta(days=2))
    client.force_login(arjuna)
    response = client.get(reverse("analysis"), {"compare": "1"})
    assert response.context["compare"]["festival_note"] == {"names": "Janmashtami", "in_current": True}
    assert "This period includes festival days (Janmashtami)" in response.content.decode()
    dashboard = client.get(reverse("dashboard", args=[arjuna.pk]), {"compare": "1"})
    assert dashboard.context["compare"]["festival_note"]["in_current"] is True

    # A festival in both periods: no note.
    Festival.objects.create(name="Ekadashi", start_date=today - timedelta(days=9), end_date=today - timedelta(days=9))
    assert client.get(reverse("analysis"), {"compare": "1"}).context["compare"]["festival_note"] is None


def test_days_filter_applies_to_both_periods(client, two_weeks, today):
    arjuna, _ = two_weeks
    Festival.objects.create(name="Ekadashi", start_date=today - timedelta(days=9), end_date=today - timedelta(days=9))
    client.force_login(arjuna)
    response = client.get(reverse("analysis"), {"compare": "1", "days": "regular"})
    total = next(item["change"] for item in response.context["compare"]["summary"] if str(item["label"]) == "Total")
    # The festival day (30 minutes) drops out of the previous week only.
    assert (total.current, total.previous) == (540, 180)
    table = {row.user.display_name: change for row, change, *_levels in response.context["table"]}
    assert table["Arjuna"].previous == 30.0  # 180 over the 6 regular days


def test_exports_include_the_comparison_only_when_switched_on(client, two_weeks):
    arjuna, _ = two_weeks
    client.force_login(arjuna)

    def workbook(name, *args, **params):
        return load_workbook(BytesIO(client.get(reverse(name, args=args), params).content))

    plain = workbook("analysis_export")
    assert "Previous daily average" not in [cell.value for row in plain.active.iter_rows() for cell in row]
    compared = [list(row) for row in workbook("analysis_export", compare="1").active.iter_rows(values_only=True)]
    header = next(row for row in compared if row[0] == "Name")
    assert header[-3:] == ["Previous daily average", "Change", "Change %"]
    arjuna_row = next(row for row in compared if row[0] == "Arjuna")
    assert arjuna_row[-3:] == [30, 30, 100]
    bhima_row = next(row for row in compared if row[0] == "Bhima")
    assert bhima_row[-3:] == [None, None, None]

    assert workbook("dashboard_export", arjuna.pk).sheetnames == ["Entries"]
    book = workbook("dashboard_export", arjuna.pk, compare="1")
    assert book.sheetnames == ["Entries", "Comparison"]
    rows = [list(row) for row in book["Comparison"].iter_rows(values_only=True)]
    assert ["Hearing", "Daily average", 60, 30, 30, 100] in rows


def test_other_filters_keep_the_comparison_on(client, two_weeks):
    arjuna, _ = two_weeks
    client.force_login(arjuna)
    # Field / Days / Group form on Analysis, and the heatmap field selector on My Report.
    assert client.get(reverse("analysis"), {"compare": "1"}).content.decode().count('name="compare" value="1"') >= 3
    report = client.get(reverse("dashboard", args=[arjuna.pk]), {"compare": "1"}).content.decode()
    assert report.count('name="compare" value="1"') >= 3
    assert 'name="compare"' not in client.get(reverse("analysis")).content.decode()
